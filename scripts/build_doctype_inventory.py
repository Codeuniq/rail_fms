"""Build an Excel inventory of the rail_fms DocTypes ported from Access."""

import json
import os
import sys

import frappe
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

APP = "/Users/sherinkr/workspace/version-15/apps/rail_fms"
OUT = "/Users/sherinkr/workspace/rail_fms_doctypes.xlsx"
sys.path.insert(0, APP)
sys.path.insert(0, os.path.join(APP, "scripts"))

import doctype_spec as spec  # noqa: E402

# DocType -> (Type, Sub Type, note). Type is the Master/Transaction split asked for;
# Sub Type carries the detail that split alone would hide.
MASTER = "Master"
TXN = "Transaction"

CLASSIFICATION = {
	# -- Rail Masters
	"Railway Station": (MASTER, "Master", "Split from Stations: parent keyed on the station code"),
	"Railway Station Category": (MASTER, "Child Table", "The per-category rows of Stations"),
	"Train": (MASTER, "Master", "6 duplicate train numbers in Access were dropped"),
	"Commodity": (MASTER, "Master", "Split from Commodities: parent keyed on the commodity name"),
	"Commodity Category": (MASTER, "Child Table", "The per-category rows of Commodities"),
	"Lease Rate": (MASTER, "Master", "Lease rate card per train, wagon and party"),
	"Special Train": (MASTER, "Master", "147 rows carry Rec_ID 0 as a 'no id' sentinel"),
	"Special Train Class": (MASTER, "Master", "Class capacity and fare per special train"),
	"Station Code Map": (MASTER, "Master", "CCRB code to console code alias"),
	"PRS Code Map": (MASTER, "Master", "CCRB code to console code alias"),
	"Train Renumbering": (MASTER, "Master", "Old to new train number, with effect date"),
	"Train Renumbering New": (MASTER, "Master", "Empty in Access"),
	"Rail FMS Settings": (MASTER, "Settings", "Single DocType: zone, division, background"),
	"Suggestion": (TXN, "Log", "Empty in Access"),
	"Application Error": (TXN, "Log", "Empty in Access"),
	# -- Passenger Earnings
	"Station Earning": (TXN, "Daily Transaction", "Live table, 2003-05 to 2023-05"),
	"Station Earning 2011": (TXN, "Archive Snapshot", "Backup copy, 2000-06 to 2014-08"),
	"Station Earning 2013": (TXN, "Archive Snapshot", "Backup copy, 2010-10 to 2013-03"),
	"Station Earning Mar 2006": (TXN, "Archive Snapshot", "Backup copy, 2000-06 to 2007-11"),
	"Station Earning Snapshot": (TXN, "Archive Snapshot", "Single day, 2005-12-01"),
	"Station Earning Copy": (TXN, "Archive Snapshot", "Backup copy, 2000-06 to 2005-12"),
	"Station Earning Staging": (TXN, "Staging", "Import staging table"),
	"PRS Earning": (TXN, "Daily Transaction", "Live table, 2016-09 to 2023-05"),
	"PRS Earning 2011": (TXN, "Archive Snapshot", "Backup copy, 2004-03 to 2014-08"),
	"PRS Earning 2013": (TXN, "Archive Snapshot", "Backup copy, 2012-01 to 2013-03"),
	"PRS Earning 2015": (TXN, "Archive Snapshot", "Backup copy, 2013-04 to 2015-03"),
	"PRS Earning Current": (TXN, "Archive Snapshot", "Backup copy, 2009-01 to 2022-06"),
	"PRS Earning Mar 2006": (TXN, "Archive Snapshot", "Backup copy, 2000-05 to 2007-11"),
	"PRS Earning 2018": (TXN, "Archive Snapshot", "Backup copy, 2017-04 to 2021-06"),
	"PRS Earning Copy 30 Jun 2021": (TXN, "Archive Snapshot", "Near-identical to PRS Earning 2018"),
	"PRS Earning Staging": (TXN, "Staging", "Import staging table, empty in Access"),
	"Division Earning": (TXN, "Daily Transaction", "Empty in Access"),
	"Comparative Summary": (TXN, "Summary", "Saved-query result snapshot"),
	"PRS Summary Current": (TXN, "Summary", "Saved-query result snapshot (Expr1001 columns)"),
	"PRS Summary Previous": (TXN, "Summary", "Saved-query result snapshot (Expr1001 columns)"),
	"Station Summary Current": (TXN, "Summary", "Saved-query result snapshot (Expr1001 columns)"),
	"Station Summary Previous": (TXN, "Summary", "Saved-query result snapshot (Expr1001 columns)"),
	"UTS On Mobile": (TXN, "Daily Transaction", "Unreserved ticketing, 2018-02 to 2020-01"),
	# -- Reservation And Occupancy
	"Waitlist Position": (TXN, "Daily Transaction", "Class-wise waitlist, 2004-03 to 2023-05"),
	"Waitlist Position Oct 2004": (TXN, "Archive Snapshot", "Subset of Waitlist Position"),
	"Special Train Occupancy": (TXN, "Daily Transaction", "Coach occupancy per special train per day"),
	"Extra Coach Attachment": (TXN, "Daily Transaction", "Extra coaches attached per train per day"),
	"Unmanned Coach Check": (TXN, "Daily Transaction", "2010-01 to 2019-04"),
	"Unmanned Coach Check Staging": (TXN, "Staging", "Empty in Access"),
	# -- Parcel And Goods
	"Parcel Daily Position": (TXN, "Daily Transaction", "54-column daily parcel MIS"),
	"Parcel Special Train": (TXN, "Daily Transaction", "Parcel special trips, 2002-08 to 2004-08"),
	"Parcel Earning": (TXN, "Daily Transaction", "Daily parcel and luggage earnings"),
	"Parcel Earning Cumulative": (TXN, "Summary", "Month-end cumulative parcel earnings"),
	"Parcel Left Over": (TXN, "Daily Transaction", "Left-over parcels by section and destination"),
	"Commodity Loading": (TXN, "Daily Transaction", "Goods loading by commodity and party"),
	"Wagon Demand": (TXN, "Daily Transaction", "Only 3 rows in Access"),
	"Wagon Release": (TXN, "Daily Transaction", "Wagon placement and release, 2001-05 to 2023-05"),
	"On Hand VPH": (TXN, "Daily Transaction", "Wagons on hand by arrival date"),
	"VPU Loading": (TXN, "Daily Transaction", "Parcel van loading, 2002-08 to 2023-05"),
	"Lease Loading": (TXN, "Daily Transaction", "Leased SLR/wagon loading, 2004-04 to 2023-05"),
	"Lease Loading New": (TXN, "Archive Snapshot", "Same 216,529 rows as Lease Loading"),
	"SLR Utilisation": (TXN, "Daily Transaction", "2001-01 to 2023-05"),
	"SLR Monthly Utilisation": (TXN, "Summary", "Monthly SLR utilisation percentage"),
	# -- Operations And Performance
	"Equipment Failure": (TXN, "Daily Transaction", "Failure windows by station, 2007-01 to 2023-04"),
	"Division Actual": (TXN, "Summary", "Month-end divisional actuals, 1995-04 to 2023-04"),
	"Division Target": (TXN, "Summary", "Month-end divisional targets, 1995-04 to 2020-03"),
}

LEGACY_NOTE = {
	"Legacy Sheet1": "Excel import staging sheet",
	"Legacy Sheet2": "Excel import staging sheet",
	"Legacy Jan 2005": "Untyped text dump, columns named 1-6",
	"Legacy Feb 2005": "Untyped text dump, columns named 1-6",
	"Legacy Mar 2005": "Untyped text dump, columns named 1-6",
	"Legacy Jun 2005": "Untyped text dump, columns named 1-6",
	"Legacy Paste Error": "Access paste-error artifact",
}

HEADERS = [
	"DocType",
	"Type",
	"Sub Type",
	"Module",
	"Access Table",
	"Access Rows",
	"Loaded Records",
	"Notes",
]


def classify(doctype):
	if doctype in CLASSIFICATION:
		return CLASSIFICATION[doctype]
	note = LEGACY_NOTE.get(doctype, "Access import/export error artifact")
	return (TXN, "Staging", note)


def build_rows(schema):
	rows = []
	for module, tables in spec.MODULES.items():
		for access_table, doctype in tables.items():
			access_rows = schema[access_table]["row_count"]
			rows.append(row_for(doctype, module, access_table, access_rows))

			split = spec.SPLIT_TABLES.get(access_table)
			if split:
				rows.append(row_for(split["child_doctype"], module, access_table, access_rows))
	return rows


def row_for(doctype, module, access_table, access_rows):
	kind, sub_type, note = classify(doctype)
	loaded = frappe.db.count(doctype) if not frappe.get_meta(doctype).issingle else ""
	return [doctype, kind, sub_type, module, access_table, access_rows, loaded, note]


def write_workbook(rows):
	workbook = Workbook()
	sheet = workbook.active
	sheet.title = "DocTypes"
	style_header(sheet, HEADERS)

	for row in sorted(rows, key=lambda r: (r[1] != MASTER, r[3], r[0])):
		sheet.append(row)

	fills = {
		MASTER: PatternFill("solid", fgColor="E8F1DE"),
		TXN: PatternFill("solid", fgColor="FDF2E3"),
	}
	for row in sheet.iter_rows(min_row=2):
		fill = fills[row[1].value]
		for cell in row[:3]:
			cell.fill = fill
		row[5].number_format = "#,##0"
		row[6].number_format = "#,##0"

	size_columns(sheet, [34, 13, 19, 27, 28, 13, 16, 58])
	sheet.freeze_panes = "A2"
	sheet.auto_filter.ref = sheet.dimensions

	write_summary(workbook, rows)
	workbook.save(OUT)


def write_summary(workbook, rows):
	sheet = workbook.create_sheet("Summary")
	style_header(sheet, ["Grouping", "Value", "DocTypes", "Access Rows"])

	groups = [("Type", 1), ("Sub Type", 2), ("Module", 3)]
	for label, index in groups:
		totals = {}
		counted = set()
		for row in rows:
			key = row[index]
			count, access_rows = totals.get(key, (0, 0))
			# Stations and Commodities each feed a parent and a child DocType. Their
			# rows are counted once, against the parent, so every grouping totals the same.
			if row[4] not in counted:
				counted.add(row[4])
				access_rows += row[5]
			totals[key] = (count + 1, access_rows)
		for key, (count, access_rows) in sorted(totals.items(), key=lambda kv: -kv[1][0]):
			sheet.append([label, key, count, access_rows])
		sheet.append([])

	sheet.append([])
	sheet.append([
		"Note",
		"Stations and Commodities each feed a parent and a child DocType. "
		"Their Access rows are counted once, against the parent.",
	])
	sheet.cell(row=sheet.max_row, column=1).font = Font(bold=True)

	for row in sheet.iter_rows(min_row=2):
		if row[3].value is not None:
			row[3].number_format = "#,##0"
	size_columns(sheet, [14, 26, 12, 14])
	sheet.freeze_panes = "A2"


def style_header(sheet, headers):
	sheet.append(headers)
	for cell in sheet[1]:
		cell.font = Font(bold=True, color="FFFFFF")
		cell.fill = PatternFill("solid", fgColor="34495E")
		cell.alignment = Alignment(vertical="center")
	sheet.row_dimensions[1].height = 22


def size_columns(sheet, widths):
	for index, width in enumerate(widths, start=1):
		sheet.column_dimensions[get_column_letter(index)].width = width


def main():
	frappe.init(site="rail_fms.dev")
	frappe.connect()
	with open(os.path.join(APP, "schema", "access_schema.json")) as handle:
		schema = json.load(handle)
	rows = build_rows(schema)
	write_workbook(rows)
	masters = sum(1 for row in rows if row[1] == MASTER)
	print(f"{len(rows)} DocTypes ({masters} master, {len(rows) - masters} transaction) -> {OUT}")
	frappe.destroy()


if __name__ == "__main__":
	main()
